"""Read the calculation workbook and draw the sensitivity curves."""
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / 'figures'


def main():
    OUT.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(ROOT / 'УИР_2_вариант_16_расчеты.xlsx', data_only=True)
    ws = wb['Сравнение']
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 9})
    assert ws['A15'].value == 'λ, 1/с'
    x = [ws.cell(row, 1).value for row in range(16, 36)]
    settings = [
        (2, 3, r'$\Lambda^{\prime}$, заявок/с', 'throughput_sensitivity.pdf'),
        (4, 5, r'$\pi$', 'loss_sensitivity.pdf'),
        (6, 7, r'$U$, с', 'sojourn_sensitivity.pdf'),
    ]
    for c1, c2, ylabel, filename in settings:
        fig, ax = plt.subplots(figsize=(7.0, 2.9))
        ax.plot(x, [ws.cell(row, c1).value for row in range(16, 36)],
                color='black', marker='o', markersize=3, linewidth=1,
                label='СИСТЕМА_1')
        ax.plot(x, [ws.cell(row, c2).value for row in range(16, 36)],
                color='0.45', marker='s', markersize=3, linewidth=1,
                linestyle='--', label='СИСТЕМА_2')
        ax.axvline(0.8, color='0.6', linestyle=':', linewidth=0.8)
        ax.set(xlabel=r'$\lambda$, 1/с', ylabel=ylabel)
        ax.grid(True, linewidth=0.35, alpha=0.45)
        ax.legend(frameon=False)
        fig.tight_layout()
        fig.savefig(OUT / filename, bbox_inches='tight')
        plt.close(fig)


if __name__ == '__main__':
    main()
